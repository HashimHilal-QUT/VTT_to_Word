import re
from io import BytesIO

import streamlit as st
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn


FONT_NAME = "Calibri"
FONT_SIZE = 12


def set_document_font(document):
    style = document.styles["Normal"]

    style.font.name = FONT_NAME
    style.font.size = Pt(FONT_SIZE)

    style.element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    style.element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)


def is_timestamp(line):
    pattern = (
        r"^\d{2}:\d{2}:\d{2}\.\d{3}\s+-->\s+"
        r"\d{2}:\d{2}:\d{2}\.\d{3}"
    )

    return re.match(pattern, line.strip()) is not None


def parse_vtt(vtt_text):
    lines = vtt_text.splitlines()

    transcript_blocks = []
    current_block = []

    for raw_line in lines:
        line = raw_line.strip()

        if not line:
            if current_block:
                transcript_blocks.append(" ".join(current_block))
                current_block = []
            continue

        if line.upper() == "WEBVTT":
            continue

        if line.isdigit():
            continue

        if is_timestamp(line):
            continue

        if line.startswith(("NOTE", "STYLE", "REGION")):
            continue

        line = re.sub(r"<[^>]+>", "", line)

        current_block.append(line)

    if current_block:
        transcript_blocks.append(" ".join(current_block))

    return transcript_blocks


def create_docx(transcript_blocks):
    document = Document()

    set_document_font(document)

    for block in transcript_blocks:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)

        match = re.match(r"^([^:]+):\s*(.*)$", block)

        if match:
            speaker = match.group(1).strip()
            dialogue = match.group(2).strip()

            speaker_run = paragraph.add_run(f"{speaker}: ")
            speaker_run.bold = True
            speaker_run.font.name = FONT_NAME
            speaker_run.font.size = Pt(FONT_SIZE)

            dialogue_run = paragraph.add_run(dialogue)
            dialogue_run.font.name = FONT_NAME
            dialogue_run.font.size = Pt(FONT_SIZE)

        else:
            run = paragraph.add_run(block)
            run.font.name = FONT_NAME
            run.font.size = Pt(FONT_SIZE)

    output = BytesIO()
    document.save(output)
    output.seek(0)

    return output


# ---------------------------------------------------------
# Streamlit UI
# ---------------------------------------------------------

st.set_page_config(
    page_title="VTT to DOCX Converter",
    page_icon="📄",
    layout="centered",
)

st.title("VTT to DOCX Transcript Converter")

st.write(
    "Upload a WebVTT (.vtt) transcript and convert it into a "
    "formatted Microsoft Word (.docx) document."
)

st.info(
    "Speaker turns are preserved. VTT timestamps, cue numbers, "
    "and metadata are removed."
)

uploaded_file = st.file_uploader(
    "Upload VTT transcript",
    type=["vtt"]
)

if uploaded_file is not None:

    vtt_text = uploaded_file.read().decode("utf-8-sig")

    transcript = parse_vtt(vtt_text)

    st.success(
        f"Transcript loaded successfully. "
        f"{len(transcript)} transcript blocks found."
    )

    docx_file = create_docx(transcript)

    output_filename = uploaded_file.name.rsplit(".", 1)[0] + ".docx"

    st.download_button(
        label="Download DOCX",
        data=docx_file,
        file_name=output_filename,
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
    )